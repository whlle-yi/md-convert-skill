#!/usr/bin/env python3
"""md_convert.py — 高质量 Markdown → Word(.docx) / PDF / HTML 转换工具。

转换管线
--------
docx :  pandoc + 程序化生成的 reference.docx（受控样式）
        → python-docx 后处理（中西文分别加粗、三线表、目录字段）
pdf  :  ① Word COM（Windows + Word，保真度最高，推荐）
        ② XeLaTeX（学术排版，公式/图表质量好）
        ③ LibreOffice（兜底）
html :  pandoc 独立单文件（内嵌 CSS，可直接分发）

排版规范（默认值，均可用命令行参数覆盖）
    中文正文 宋体；西文/数字 Times New Roman；代码 Consolas
    中文标题与加粗 思源宋体 Heavy（不使用伪粗体）
    正文 12pt（小四）、1.5 倍行距、首行缩进 2 字符
    表格默认三线表；公式转 Word 原生公式（OMML）/ LaTeX 公式

用法示例
    python md_convert.py report.md -o report.docx
    python md_convert.py report.md -o report.pdf
    python md_convert.py report.md -o report.html --toc
    python md_convert.py paper.md -o paper.pdf --pdf-engine latex --number-sections
"""

from __future__ import annotations

import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_reference_docx import build_reference  # noqa: E402

__version__ = "0.1.0"

# CJK 统一表意文字及扩展、兼容表意、部首、中文标点、全角字符
CJK_RE = re.compile(
    "[\u2e80-\u2eff\u3000-\u303f\u31c0-\u31ef\u3200-\u32ff\u3400-\u4dbf"
    "\u4e00-\u9fff\uf900-\ufaff\ufe30-\ufe4f\uff00-\uffef]"
)
SEG_RE = re.compile("([\u2e80-\u2eff\u3000-\u303f\u31c0-\u31ef\u3200-\u32ff\u3400-\u4dbf"
                    "\u4e00-\u9fff\uf900-\ufaff\ufe30-\ufe4f\uff00-\uffef]+"
                    "|[^\u2e80-\u2eff\u3000-\u303f\u31c0-\u31ef\u3200-\u32ff\u3400-\u4dbf"
                    "\u4e00-\u9fff\uf900-\ufaff\ufe30-\ufe4f\uff00-\uffef]+)")

DEFAULT_FROM = "markdown+east_asian_line_breaks"
DEFAULT_HEAVY_FONT = "Noto Serif SC Heavy X"
DEFAULT_HEAVY_FONT_FILE = Path.home() / ".zcode" / "fonts" / "NotoSerifSC-HeavyX.ttf"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# OOXML 子元素顺序（CT_* 序列的后继片段），插入元素时必须遵守，
# 否则 Word 会判定文档无效并进入修复/只读模式。
_RPR_AFTER_RFONTS = (
    "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike",
    "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
    "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
    "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
    "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
)
_RPR_AFTER_B = _RPR_AFTER_RFONTS[2:]
_TBLPR_AFTER_TBLW = (
    "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
)
_TBLPR_AFTER_BORDERS = _TBLPR_AFTER_TBLW[4:]
_TRPR_AFTER_CANTSPLIT = ("w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden")
_TRPR_AFTER_TBLHEADER = ("w:tblCellSpacing", "w:jc", "w:hidden")
_TCPR_AFTER_TCBORDERS = (
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
)


def _get_or_add(parent, tag: str, successors: tuple[str, ...]):
    """取（或按 schema 顺序创建）parent 的直接子元素 tag。"""
    el = parent.find(f"{W_NS}{tag}")
    if el is None:
        el = parent.makeelement(f"{W_NS}{tag}", {})
        parent.insert_element_before(el, *successors)
    return el


class ConvertError(RuntimeError):
    """带用户可读信息的转换错误。"""


def log(msg: str) -> None:
    print(f"[md-convert] {msg}", file=sys.stderr)


# ---------------------------------------------------------------------------
# pandoc
# ---------------------------------------------------------------------------

def require_pandoc() -> str:
    path = shutil.which("pandoc")
    if path is None:
        raise ConvertError(
            "未找到 pandoc。请安装后重试（见 references/troubleshooting.md）：\n"
            "  Windows: winget install JohnMacFarlane.Pandoc  或到 https://pandoc.org/installing.html 下载\n"
            "  macOS:   brew install pandoc\n"
            "  Linux:   sudo apt install pandoc"
        )
    return path


def run_pandoc(args: list[str], verbose: bool = False) -> None:
    cmd = ["pandoc", *args]
    if verbose:
        log("执行: " + " ".join(f'"{a}"' if " " in a else a for a in cmd))
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if proc.returncode != 0:
        raise ConvertError(f"pandoc 转换失败：\n{proc.stderr.strip() or proc.stdout.strip()}")
    if verbose and proc.stderr.strip():
        log("pandoc: " + proc.stderr.strip())


# ---------------------------------------------------------------------------
# docx 后处理（python-docx + OOXML）
# ---------------------------------------------------------------------------

def _iter_body_paragraphs(doc):
    body = doc.element.body
    for p in body.iter(f"{W_NS}p"):
        yield p


def _pstyle_of(p_el) -> str | None:
    ppr = p_el.find(f"{W_NS}pPr")
    if ppr is None:
        return None
    pstyle = ppr.find(f"{W_NS}pStyle")
    return pstyle.get(f"{W_NS}val") if pstyle is not None else None


def _effective_bold(pstyle: str | None, rpr) -> bool:
    """判断 run 是否“实际加粗”：直接 w:b > rStyle(Strong) > 段落样式(标题族)。"""
    if rpr is not None:
        b = rpr.find(f"{W_NS}b")
        if b is not None:
            return b.get(f"{W_NS}val") not in ("0", "false", "none")
        rstyle = rpr.find(f"{W_NS}rStyle")
        if rstyle is not None and (rstyle.get(f"{W_NS}val") or "") == "Strong":
            return True
    if pstyle and (pstyle.startswith("Heading") or pstyle in ("Title", "Subtitle")):
        return True
    return False


def _run_text(r_el) -> str | None:
    """返回 run 的纯文本；若包含 w:t 以外的内容（图、换行、域等）返回 None。"""
    parts: list[str] = []
    for child in r_el:
        tag = child.tag
        if tag == f"{W_NS}rPr":
            continue
        if tag == f"{W_NS}t":
            parts.append(child.text or "")
        else:
            return None
    return "".join(parts)


def _make_segment_run(r_el, text: str, *, heavy_font: str | None) -> object:
    """复制 run 并写入单段文本；heavy_font 不为 None 时按“中文强调”规则处理。"""
    new_r = copy.deepcopy(r_el)
    # 删除全部 w:t，写入新文本
    for t in new_r.findall(f"{W_NS}t"):
        new_r.remove(t)
    t = new_r.makeelement(f"{W_NS}t", {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    # w:t 必须位于 rPr 之后
    rpr = new_r.find(f"{W_NS}rPr")
    if rpr is not None:
        rpr.addnext(t)
    else:
        new_r.insert(0, t)

    if heavy_font is not None:
        rpr = new_r.find(f"{W_NS}rPr")
        if rpr is None:
            rpr = new_r.makeelement(f"{W_NS}rPr", {})
            new_r.insert(0, rpr)
        # 显式关闭加粗（run 级 w:b=0 覆盖样式级 w:b，避免 Heavy 字体再被伪粗体）
        b = _get_or_add(rpr, "b", _RPR_AFTER_B)
        b.set(f"{W_NS}val", "0")
        rfonts = _get_or_add(rpr, "rFonts", _RPR_AFTER_RFONTS)
        rfonts.set(f"{W_NS}eastAsia", heavy_font)
    return new_r


def fix_cjk_bold(doc, heavy_font: str) -> int:
    """中文强调规范化：加粗文本按中西文拆分 run。

    中文段使用 Heavy 宋体并显式关闭 w:b（不产生伪粗体）；
    西文段保留原样式（Times New Roman Bold）。
    """
    changed = 0
    for p in _iter_body_paragraphs(doc):
        pstyle = _pstyle_of(p)
        for r in list(p.iter(f"{W_NS}r")):
            text = _run_text(r)
            if not text or not CJK_RE.search(text):
                continue
            rpr = r.find(f"{W_NS}rPr")
            if not _effective_bold(pstyle, rpr):
                continue
            segs = [m.group(0) for m in SEG_RE.finditer(text)]
            if len(segs) <= 1 and CJK_RE.fullmatch(text):
                segs = [text]  # 纯中文整段，也统一改写
            anchors = []
            for seg in segs:
                is_cjk = bool(CJK_RE.match(seg))
                anchors.append(_make_segment_run(r, seg, heavy_font=heavy_font if is_cjk else None))
            for el in anchors:
                r.addprevious(el)
            p.remove(r)
            changed += 1
    return changed


def _table_autofit(tbl) -> None:
    """列宽按内容自适应。

    pandoc 依据 md 源中分隔行的横线长度分配列宽，窄数字列会被压出折行
    （如 1,024 断成 1,02/4）。清除显式列宽并切到 autofit，由 Word 按内容重排。
    """
    tblpr = tbl._tbl.tblPr
    tblw = _get_or_add(tblpr, "tblW", _TBLPR_AFTER_TBLW)
    tblw.set(f"{W_NS}w", "0")
    tblw.set(f"{W_NS}type", "auto")
    layout = _get_or_add(tblpr, "tblLayout", _TBLPR_AFTER_TBLW)
    layout.set(f"{W_NS}type", "autofit")
    for tcpr in tbl._tbl.iter(f"{W_NS}tcPr"):
        for tcw in tcpr.findall(f"{W_NS}tcW"):
            tcw.set(f"{W_NS}w", "0")
            tcw.set(f"{W_NS}type", "auto")


def _repeat_header_row(row) -> None:
    """表头行跨页重复，且行内不断页。"""
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(f"{W_NS}cantSplit") is None:
        trpr.insert_element_before(trpr.makeelement(f"{W_NS}cantSplit", {}),
                                   *_TRPR_AFTER_CANTSPLIT)
    if trpr.find(f"{W_NS}tblHeader") is None:
        trpr.insert_element_before(trpr.makeelement(f"{W_NS}tblHeader", {}),
                                   *_TRPR_AFTER_TBLHEADER)


def style_tables(doc, mode: str = "threeline") -> int:
    """为所有表格设置边框：threeline（学术三线表）或 grid（全框线）。

    同时居中表格、列宽按内容自适应、表头行加粗并跨页重复。
    表头中文随后由 fix_cjk_bold 规范化字重。
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT

    count = 0
    for tbl in doc.tables:
        tblpr = tbl._tbl.tblPr
        old = tblpr.find(f"{W_NS}tblBorders")
        if old is not None:
            tblpr.remove(old)
        if mode == "threeline":
            spec = {"top": ("single", 12), "bottom": ("single", 12),
                    "left": ("none", 0), "right": ("none", 0),
                    "insideH": ("none", 0), "insideV": ("none", 0)}
        else:
            spec = {s: ("single", 4) for s in ("top", "left", "bottom", "right", "insideH", "insideV")}
        borders = tblpr.makeelement(f"{W_NS}tblBorders", {})
        tblpr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)
        for side, (val, sz) in spec.items():
            side_el = borders.makeelement(f"{W_NS}{side}", {})
            side_el.set(f"{W_NS}val", val)
            if val == "single":
                side_el.set(f"{W_NS}sz", str(sz))
                side_el.set(f"{W_NS}space", "0")
                side_el.set(f"{W_NS}color", "000000")
            borders.append(side_el)

        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_autofit(tbl)

        # 表头行：加粗 + 三线表中的 0.75pt 下边线 + 跨页重复
        if tbl.rows:
            _repeat_header_row(tbl.rows[0])
            for cell in tbl.rows[0].cells:
                tcpr = cell._tc.get_or_add_tcPr()
                old_tc = tcpr.find(f"{W_NS}tcBorders")
                if old_tc is not None:
                    tcpr.remove(old_tc)
                if mode == "threeline":
                    tcb = tcpr.makeelement(f"{W_NS}tcBorders", {})
                    bottom = tcb.makeelement(f"{W_NS}bottom", {})
                    bottom.set(f"{W_NS}val", "single")
                    bottom.set(f"{W_NS}sz", "6")
                    bottom.set(f"{W_NS}space", "0")
                    bottom.set(f"{W_NS}color", "000000")
                    tcb.append(bottom)
                    tcpr.insert_element_before(tcb, *_TCPR_AFTER_TCBORDERS)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
        count += 1
    return count


def enable_update_fields(doc) -> None:
    """标记文档打开时更新域（目录页码在 Word 中自动刷新）。"""
    settings = doc.settings.element
    if settings.find(f"{W_NS}updateFields") is None:
        el = settings.makeelement(f"{W_NS}updateFields", {})
        el.set(f"{W_NS}val", "true")
        settings.append(el)


# ---------------------------------------------------------------------------
# 字体内嵌（ODTTF）：把 Heavy 中文字体子集嵌入 docx
# ---------------------------------------------------------------------------

def _subset_font(font_path: Path, text: str) -> bytes:
    """按文档用字裁剪字体，控制体积（完整 CJK 字体 10MB+，子集通常 <2MB）。"""
    import io

    from fontTools import subset
    from fontTools.ttLib import TTFont

    opts = subset.Options()
    opts.name_IDs = ["*"]
    opts.name_legacy = True
    opts.name_languages = ["*"]
    opts.glyph_names = False
    opts.layout_features = ["*"]
    opts.notdef_outline = True
    font = TTFont(str(font_path))
    subsetter = subset.Subsetter(options=opts)
    subsetter.populate(text=text)
    subsetter.subset(font)
    buf = io.BytesIO()
    font.save(buf)
    return buf.getvalue()


def _obfuscate_odttf(data: bytes, guid_str: str) -> bytes:
    """ODTTF 混淆：前 32 字节与 GUID 十六进制字节顺序异或，循环两轮。"""
    hexstr = guid_str.strip("{}").replace("-", "")
    key = bytes.fromhex(hexstr)
    head = bytearray(data[:32])
    for i in range(32):
        head[i] ^= key[i % 16]
    return bytes(head) + data[32:]


def embed_heavy_font(docx_path: Path, family: str, font_file: Path) -> bool:
    """把 Heavy 字体以 ODTTF 子集嵌入 docx。

    Word 对按用户注册的字体解析不稳定、对 CJK 变量字体支持差，接收方也
    往往没有该字体——内嵌后渲染不依赖系统安装。成功返回 True。
    """
    import io
    import uuid
    import zipfile

    if not font_file.is_file():
        return False
    with zipfile.ZipFile(docx_path) as z:
        names = z.namelist()
        texts = []
        for part in ("word/document.xml", "word/footnotes.xml", "word/endnotes.xml",
                     "word/header1.xml", "word/footer1.xml"):
            if part in names:
                xml = z.read(part).decode("utf-8", errors="replace")
                texts.append("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)))
        chars = "".join(texts) + family
        subset = _subset_font(font_file, chars)

    guid_hex = uuid.uuid4().hex
    guid_str = (f"{{{guid_hex[:8].upper()}-{guid_hex[8:12].upper()}-"
                f"{guid_hex[12:16].upper()}-{guid_hex[16:20].upper()}-"
                f"{guid_hex[20:32].upper()}}}")
    odttf = _obfuscate_odttf(subset, guid_str)

    with zipfile.ZipFile(docx_path) as z:
        items = {n: z.read(n) for n in z.namelist()}

    # 1) [Content_Types].xml：注册 odttf 类型
    ct = items["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace(
            "</Types>",
            '<Default Extension="odttf" '
            'ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/></Types>')
    items["[Content_Types].xml"] = ct.encode("utf-8")

    # 2) fontTable.xml：补/改 w:font 条目，挂 embedRegular
    ft = items["word/fontTable.xml"].decode("utf-8")
    font_block = (f'<w:font w:name="{family}">'
                  f'<w:charset w:val="86"/><w:family w:val="roman"/>'
                  f'<w:pitch w:val="variable"/>'
                  f'<w:embedRegular r:id="rIdMdfont1" w:fontKey="{guid_str}"/>'
                  f'</w:font>')
    if f'w:name="{family}"' in ft:
        ft = re.sub(rf'<w:font w:name="{re.escape(family)}">.*?</w:font>',
                    font_block, ft, flags=re.S)
    else:
        ft = ft.replace("</w:fonts>", font_block + "</w:fonts>")
    items["word/fontTable.xml"] = ft.encode("utf-8")

    # 3) fontTable 的 rels
    rel_name = "word/_rels/fontTable.xml.rels"
    rel_entry = ('<Relationship Id="rIdMdfont1" Type="http://schemas.openxmlformats.org/'
                 'officeDocument/2006/relationships/font" Target="fonts/font1.odttf"/>')
    if rel_name in items:
        rels = items[rel_name].decode("utf-8").replace(
            "</Relationships>", rel_entry + "</Relationships>")
    else:
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + rel_entry + "</Relationships>")
    items[rel_name] = rels.encode("utf-8")

    # 4) settings.xml：声明文档含内嵌字体（按 schema 序列插在靠前位置）
    st = items["word/settings.xml"].decode("utf-8")
    if "embedTrueTypeFonts" not in st:
        el = "<w:embedTrueTypeFonts/><w:saveSubsetFonts/>"
        m = re.search(r"<w:settings[^>]*>", st)
        open_end = m.end()
        # 找到第一个序位在 embedTrueTypeFonts 之后的子元素，插到它前面
        succ = re.search(r"<w:(?:defaultTabStop|autoHyphenation|compatability|"
                         r"compat|rsids|themeFontLang|clrSchemeMapping|shapeDefaults)"
                         r"[ />]", st[open_end:])
        if succ:
            pos = open_end + succ.start()
            st = st[:pos] + el + st[pos:]
        else:
            st = st[:open_end] + el + st[open_end:]
    items["word/settings.xml"] = st.encode("utf-8")

    # 5) 字体数据
    items["word/fonts/font1.odttf"] = odttf

    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items():
            z.writestr(name, data)
    tmp.replace(docx_path)
    return True


# ---------------------------------------------------------------------------
# 各格式转换
# ---------------------------------------------------------------------------

def convert_docx(src: Path, out: Path, args, tmpdir: Path) -> Path:
    ref = build_reference(
        tmpdir / "reference.docx",
        cjk_font=args.cjk_font,
        latin_font=args.latin_font,
        heavy_cjk_font=None if args.no_heavy else args.heading_cjk_font,
        mono_font=args.mono_font,
        body_pt=args.font_size,
        line_spacing=args.line_spacing,
        body_indent=args.body_indent,
        lang=args.lang,
    )
    pandoc_args = [
        str(src), "-f", args.reader, "-t", "docx",
        "--reference-doc", str(ref),
        "--resource-path", str(src.parent),
        "--dpi", str(args.dpi),
        "-M", f"lang={args.lang}",
        "-o", str(out),
    ]
    if not args.highlight:
        pandoc_args.append("--no-highlight")
    if args.toc:
        pandoc_args += ["--toc", "--toc-depth", str(args.toc_depth),
                        "-M", "toc-title=目录"]
    if args.number_sections:
        pandoc_args += ["--number-sections"]
    run_pandoc(pandoc_args, args.verbose)

    # ---- 后处理 ----
    from docx import Document

    doc = Document(str(out))
    n_tables = style_tables(doc, args.table_style)
    n_runs = 0
    if not args.no_heavy:
        n_runs = fix_cjk_bold(doc, args.heading_cjk_font)
    if args.toc:
        enable_update_fields(doc)
    doc.save(str(out))

    # ---- Heavy 字体内嵌（不依赖收件人安装） ----
    embedded = False
    if not args.no_heavy and args.embed_font:
        font_file = Path(args.emph_font_file) if args.emph_font_file else None
        if font_file and font_file.is_file():
            try:
                embedded = embed_heavy_font(out, args.heading_cjk_font, font_file)
            except Exception as e:  # 内嵌失败不阻断转换，回退为依赖本机字体
                log(f"字体内嵌失败（忽略，回退本机字体）：{e}")

    log(f"docx 完成：{out.name}（表格 {n_tables}，规范化的强调 run {n_runs}，"
        f"样式 {'三线表' if args.table_style == 'threeline' else '全框线'}"
        f"{'，字体内嵌' if embedded else ''}）")
    return out


def build_latex_header(args) -> str:
    """动态生成 XeLaTeX 导言（与 docx 排版规范一致）。"""
    lines: list[str] = []
    lines.append(f"\\setmainfont{{{args.latin_font}}}")
    lines.append("\\usepackage{xeCJK}")
    emph_file = Path(args.emph_font_file) if args.emph_font_file else None
    if args.no_heavy or emph_file is None or not emph_file.is_file():
        lines.append(f"\\setCJKmainfont{{{args.cjk_font}}}[AutoFakeBold=2.5]")
        lines.append(
            f"\\IfFontExistsTF{{{args.heading_cjk_font}}}"
            f"{{\\setCJKfamilyfont{{emphsong}}{{{args.heading_cjk_font}}}}}"
            f"{{\\setCJKfamilyfont{{emphsong}}{{{args.cjk_font}}}[AutoFakeBold=2.5]}}"
        )
    else:
        font_dir = emph_file.parent.as_posix()
        font_name = emph_file.name
        lines.append(
            f"\\setCJKmainfont{{{args.cjk_font}}}"
            f"[Path={font_dir}/, BoldFont={font_name}, AutoFakeBold=2.5]"
        )
        lines.append(
            f"\\setCJKfamilyfont{{emphsong}}{{{font_name}}}[Path={font_dir}/, BoldFont={font_name}]"
        )
    lines.append("\\newcommand{\\key}[1]{{\\CJKfamily{emphsong}\\bfseries #1}}")
    return "\n".join(lines) + "\n"


def convert_pdf(src: Path, out: Path, args, tmpdir: Path) -> Path:
    engine = args.pdf_engine
    if engine == "auto":
        if sys.platform == "win32":
            engine = "word"
        elif shutil.which("xelatex"):
            engine = "latex"
        elif shutil.which("soffice"):
            engine = "libreoffice"
        else:
            raise ConvertError(
                "未找到可用的 PDF 引擎。三选一：\n"
                "  1. Windows + Microsoft Word（推荐，自动使用）\n"
                "  2. XeLaTeX（TeX Live/MiKTeX）\n"
                "  3. LibreOffice（soffice 加入 PATH）\n"
                "或用 --pdf-engine 手动指定。"
            )

    if engine == "word":
        tmp_docx = tmpdir / (src.stem + ".docx")
        docx_args = copy.copy(args)
        convert_docx(src, tmp_docx, docx_args, tmpdir)
        ps1 = Path(__file__).resolve().parent / "docx2pdf.ps1"
        if ps1.is_file():
            cmd = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass",
                   "-File", str(ps1), "-InPath", str(tmp_docx.resolve()),
                   "-OutPath", str(out.resolve())]
        else:
            try:
                import docx2pdf  # noqa: F401
            except ImportError:
                raise ConvertError("缺少 docx2pdf.ps1 且未安装 docx2pdf，无法用 Word 转 PDF。")
            cmd = [sys.executable, "-c",
                   f"import docx2pdf; docx2pdf.convert(r'{tmp_docx}', r'{out}')"]
        if args.verbose:
            log("执行: " + " ".join(cmd))
        proc = subprocess.run(cmd, capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=300)
        if proc.returncode != 0 or not out.exists():
            detail = (proc.stderr or proc.stdout or "").strip()
            raise ConvertError(f"Word COM 转 PDF 失败：{detail}")
        log(f"pdf 完成（Word 引擎）：{out.name}")
        return out

    if engine == "latex":
        if not shutil.which("xelatex"):
            raise ConvertError("未找到 xelatex。请安装 TeX Live / MiKTeX，或改用 --pdf-engine word。")
        header = tmpdir / "header.tex"
        header.write_text(build_latex_header(args), encoding="utf-8")
        pandoc_args = [
            str(src), "-f", args.reader,
            "--pdf-engine", "xelatex",
            "-H", str(header),
            "-V", "geometry:top=2.54cm", "-V", "geometry:bottom=2.54cm",
            "-V", "geometry:left=3.17cm", "-V", "geometry:right=3.17cm",
            "-V", f"linestretch={args.line_spacing}",
            "-V", "colorlinks=true", "-V", "linkcolor=black",
            "-V", "urlcolor=black", "-V", "citecolor=black",
            "--resource-path", str(src.parent),
            "--dpi", str(args.dpi),
            "-M", f"lang={args.lang}",
            "-o", str(out),
        ]
        if args.font_size in (10.0, 11.0, 12.0):
            pandoc_args += ["-V", f"fontsize={int(args.font_size)}pt"]
        if not args.highlight:
            pandoc_args.append("--no-highlight")
        if args.toc:
            pandoc_args += ["--toc", "--toc-depth", str(args.toc_depth),
                            "-M", "toc-title=目录"]
        if args.number_sections:
            pandoc_args += ["--number-sections"]
        run_pandoc(pandoc_args, args.verbose)
        log(f"pdf 完成（XeLaTeX 引擎）：{out.name}")
        return out

    if engine == "libreoffice":
        soffice = shutil.which("soffice")
        if not soffice:
            raise ConvertError("未找到 soffice（LibreOffice）。")
        tmp_docx = tmpdir / (src.stem + ".docx")
        convert_docx(src, tmp_docx, copy.copy(args), tmpdir)
        cmd = [soffice, "--headless", "--convert-to", "pdf",
               "--outdir", str(out.parent.resolve()), str(tmp_docx.resolve())]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        produced = out.parent / (tmp_docx.stem + ".pdf")
        if proc.returncode != 0 or not produced.exists():
            raise ConvertError(f"LibreOffice 转 PDF 失败：{(proc.stderr or '').strip()}")
        shutil.move(str(produced), str(out))
        log(f"pdf 完成（LibreOffice 引擎）：{out.name}")
        return out

    raise ConvertError(f"未知 PDF 引擎：{engine}（可选 auto/word/latex/libreoffice）")


def extract_title(src: Path) -> str:
    """取 YAML front matter 的 title，否则第一个一级标题，否则文件名。"""
    try:
        text = src.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = src.read_text(encoding="utf-16")
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if m:
        tm = re.search(r'^title:\s*["\']?(.+?)["\']?\s*$', m.group(1), re.M)
        if tm:
            return tm.group(1)
    fence = False
    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            fence = not fence
            continue
        if not fence:
            hm = re.match(r"^#\s+(.+?)\s*#*\s*$", line)
            if hm:
                return hm.group(1)
    return src.stem


def convert_html(src: Path, out: Path, args, tmpdir: Path) -> Path:
    css = Path(__file__).resolve().parent.parent / "assets" / "html-style.css"
    pandoc_args = [
        str(src), "-f", args.reader, "-t", "html5",
        "--standalone", "--embed-resources", "--mathml",
        "-M", f"title={extract_title(src)}",
        "-M", f"lang={args.lang}",
        "--resource-path", str(src.parent),
        "--dpi", str(args.dpi),
        "-o", str(out),
    ]
    if css.is_file():
        pandoc_args += ["-c", str(css.resolve())]
    if not args.highlight:
        pandoc_args.append("--no-highlight")
    if args.toc:
        pandoc_args += ["--toc", "--toc-depth", str(args.toc_depth)]
    if args.number_sections:
        pandoc_args += ["--number-sections"]
    run_pandoc(pandoc_args, args.verbose)
    log(f"html 完成：{out.name}")
    return out


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="md_convert.py",
        description="高质量 Markdown → Word / PDF / HTML 转换（md-convert）",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("input", help="输入 Markdown 文件")
    p.add_argument("-o", "--output",
                   help="输出文件；--to 多产出时为输出目录（默认与输入同目录）")
    p.add_argument("-f", "--format", choices=("docx", "pdf", "html"),
                   help="目标格式（默认 docx，或按 -o 后缀推断）")
    p.add_argument("--to",
                   help="一次产出多种格式，逗号分隔（如 docx,pdf,html 或 all）；与 -f 二选一")
    p.add_argument("--reader", default=DEFAULT_FROM,
                   help=f"pandoc 输入格式（默认 {DEFAULT_FROM}；纯 GFM 用 gfm）")
    p.add_argument("--toc", action="store_true", help="生成目录（Word 中打开后自动刷新页码）")
    p.add_argument("--toc-depth", type=int, default=3, help="目录深度")
    p.add_argument("--number-sections", action="store_true", help="标题编号")
    p.add_argument("--cjk-font", default="SimSun", help="中文正文字体")
    p.add_argument("--latin-font", default="Times New Roman", help="西文/数字字体")
    p.add_argument("--mono-font", default="Consolas", help="代码字体")
    p.add_argument("--heading-cjk-font", default=DEFAULT_HEAVY_FONT,
                   help=f"中文标题/加粗字体（默认 {DEFAULT_HEAVY_FONT}）")
    p.add_argument("--no-heavy", action="store_true",
                   help="不使用 Heavy 字体，标题/加粗回退为正文字体伪粗体")
    p.add_argument("--no-embed-font", dest="embed_font", action="store_false",
                   help="关闭字体内嵌（默认内嵌 Heavy 字体子集，接收方无需安装）")
    p.add_argument("--emph-font-file", default=str(DEFAULT_HEAVY_FONT_FILE),
                   help="Heavy 字体文件路径（供 XeLaTeX 嵌入；auto 检测失败时需显式给出）")
    p.add_argument("--font-size", type=float, default=12.0, help="正文字号（pt）")
    p.add_argument("--line-spacing", type=float, default=1.5, help="正文行距")
    p.add_argument("--no-body-indent", dest="body_indent", action="store_false",
                   help="关闭正文首行缩进 2 字符")
    p.add_argument("--table-style", choices=("threeline", "grid"), default="threeline",
                   help="表格样式")
    p.add_argument("--highlight", action="store_true",
                   help="代码语法高亮（默认关闭，保持黑白正式文档风格）")
    p.add_argument("--pdf-engine", choices=("auto", "word", "latex", "libreoffice"),
                   default="auto", help="PDF 引擎")
    p.add_argument("--dpi", type=int, default=96, help="图片分辨率换算 dpi")
    p.add_argument("--lang", default="zh-CN", help="文档语言（影响断行/校对）")
    p.add_argument("--verbose", action="store_true", help="输出详细日志")
    p.add_argument("--keep-temp", action="store_true", help="保留临时目录（调试用）")
    p.add_argument("--version", action="version", version=f"md-convert {__version__}")
    return p


def _plan_outputs(args, src: Path) -> dict[str, Path] | None:
    """解析目标格式集合与各输出路径；出错打印信息并返回 None。"""
    if args.to and args.format:
        print("[md-convert] 错误：--to 与 -f/--format 二选一", file=sys.stderr)
        return None
    if args.to:
        fmts = [f.strip().lower() for f in args.to.split(",") if f.strip()]
        if fmts == ["all"]:
            fmts = ["docx", "pdf", "html"]
        bad = [f for f in fmts if f not in ("docx", "pdf", "html")]
        if bad:
            print(f"[md-convert] 错误：--to 含未知格式 {bad}，支持 docx / pdf / html", file=sys.stderr)
            return None
        fmts = list(dict.fromkeys(fmts))  # 去重保序
        if args.output:
            out_base = Path(args.output)
            if out_base.suffix:
                print("[md-convert] 错误：--to 多产出时 -o 应为目录（不带扩展名）", file=sys.stderr)
                return None
        else:
            out_base = src.parent
        return {f: out_base / (src.stem + "." + f) for f in fmts}

    fmt = args.format
    if args.output:
        out = Path(args.output)
        if fmt is None:
            fmt = out.suffix.lstrip(".").lower()
    else:
        fmt = fmt or "docx"
        out = src.with_suffix("." + fmt)
    if fmt not in ("docx", "pdf", "html"):
        print(f"[md-convert] 错误：无法识别的输出格式“{fmt}”，支持 docx / pdf / html", file=sys.stderr)
        return None
    return {fmt: out}


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass
    args = build_arg_parser().parse_args(argv)

    src = Path(args.input)
    if not src.is_file():
        print(f"[md-convert] 错误：输入文件不存在：{src}", file=sys.stderr)
        return 2

    outs = _plan_outputs(args, src)
    if outs is None:
        return 2
    for out in outs.values():
        out.parent.mkdir(parents=True, exist_ok=True)
        if out.exists():
            try:
                with open(out, "r+b"):
                    pass
            except PermissionError:
                print(f"[md-convert] 错误：输出文件被其他程序占用，无法覆盖：{out}\n"
                      "  请关闭正在预览/编辑该文件的程序（PDF 阅读器、Word、预览窗格等）后重试。",
                      file=sys.stderr)
                return 3

    require_pandoc()
    tmpdir = Path(tempfile.mkdtemp(prefix="md-convert-"))
    failed = []
    try:
        for fmt, out in outs.items():
            try:
                if fmt == "docx":
                    convert_docx(src, out, args, tmpdir)
                elif fmt == "pdf":
                    convert_pdf(src, out, args, tmpdir)
                else:
                    convert_html(src, out, args, tmpdir)
            except ConvertError as e:
                print(f"[md-convert] 错误：{e}", file=sys.stderr)
                failed.append(fmt)
    finally:
        if args.keep_temp:
            log(f"临时目录已保留：{tmpdir}")
        else:
            shutil.rmtree(tmpdir, ignore_errors=True)
    if failed:
        log(f"{len(failed)} 种格式转换失败：{'、'.join(failed)}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
